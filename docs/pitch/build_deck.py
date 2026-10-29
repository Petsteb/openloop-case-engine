"""Build docs/pitch/OpenLoop-deck.pptx from the content in docs/pitch/deck.md.

python docs/pitch/build_deck.py   (needs python-pptx). Palette = the app's "Night ledger": dark, colour only for meaning.
"""
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

OUT = Path(__file__).with_name("OpenLoop-deck.pptx")
BG, SURFACE, LINE = "111317", "1C1F26", "2E333D"
INK, MUTED = "EEF0F3", "A3AAB5"
ATTN, CONFLICT, REVIEW, EXPIRED, OK = "5B86F7", "EF5B52", "E8A93C", "A08CF7", "43B884"
FONT = "Calibri"
W, H = 13.333, 7.5

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(W), Inches(H)
BLANK = prs.slide_layouts[6]


def rgb(hex_):
    return RGBColor.from_string(hex_)


def new_slide(notes: str = ""):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb(BG)
    if notes:
        s.notes_slide.notes_text_frame.text = notes
    return s


def text(s, x, y, w, h, runs, size=16, color=INK, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         spacing=1.1, after=0):
    """runs: a string, or a list of paragraphs; a paragraph is a string or a list of (text, {opts}) runs."""
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    paras = runs if isinstance(runs, list) else [runs]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment, p.line_spacing, p.space_after = align, spacing, Pt(after)
        for chunk, opts in ([(para, {})] if isinstance(para, str) else para):
            r = p.add_run()
            r.text = chunk
            f = r.font
            f.name, f.size = FONT, Pt(opts.get("size", size))
            f.bold = opts.get("bold", bold)
            f.italic = opts.get("italic", False)
            f.color.rgb = rgb(opts.get("color", color))
    return tb


def box(s, x, y, w, h, fill=SURFACE, line=None, radius=0.06, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    sh = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        sh.adjustments[0] = radius
    sh.fill.solid()
    sh.fill.fore_color.rgb = rgb(fill)
    if line:
        sh.line.color.rgb = rgb(line)
        sh.line.width = Pt(1)
    else:
        sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def dot(s, x, y, d, color, ring=False):
    sh = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    sh.shadow.inherit = False
    if ring:
        sh.fill.background()
        sh.line.color.rgb = rgb(color)
        sh.line.width = Pt(2)
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = rgb(color)
        sh.line.fill.background()
    return sh


def arrow(s, x1, y1, x2, y2, color=MUTED):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(1.5)
    ln = c.line._get_or_add_ln()
    tail = etree.SubElement(ln, "{http://schemas.openxmlformats.org/drawingml/2006/main}tailEnd")
    tail.set("type", "triangle")
    return c


def title(s, t, kicker=None):
    text(s, 0.6, 0.45, 12.1, 0.8, t, size=36, bold=True)
    if kicker:
        text(s, 0.6, 1.25, 12.1, 0.5, kicker, size=16, color=MUTED)


def ledger_motif(s, y, x0=0.6, n=14, gap=0.42, colors=None):
    """The app's timeline: a row of event dots on a hairline (the visual motif of the deck)."""
    line = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x0), Inches(y + 0.06),
                                  Inches(x0 + gap * (n - 1) + 0.12), Inches(y + 0.06))
    line.line.color.rgb, line.line.width = rgb(LINE), Pt(1)
    colors = colors or {}
    for i in range(n):
        c = colors.get(i, MUTED)
        dot(s, x0 + i * gap, y, 0.12, c, ring=(c == MUTED and i % 3 == 1))


def stat(s, x, y, big, small, color=INK, w=3.6):
    text(s, x, y, w, 0.95, big, size=54, bold=True, color=color)
    text(s, x, y + 0.98, w, 0.8, small, size=14, color=MUTED, spacing=1.15)


# ---------------------------------------------------------------- 1. title
s = new_slide("Hook: every bank has AI that reads a document; nobody owns the file between documents. OpenLoop does.")
ledger_motif(s, 1.2, colors={3: ATTN, 6: CONFLICT, 8: REVIEW, 10: EXPIRED, 12: OK})
text(s, 0.6, 2.2, 12, 1.4, "OpenLoop", size=80, bold=True)
text(s, 0.6, 3.65, 11.5, 1.0, "The AI case agent for bank files that stay open for weeks", size=28, color=INK)
text(s, 0.6, 4.6, 11.5, 0.6, "Code enforces the rules · Claude plans the next step · people make the decisions", size=18,
     color=MUTED)
text(s, 0.6, 6.55, 12, 0.4, "NEXXT AI Hackathon 2026 · Raiffeisen Bank Romania · AI for Operational Efficiency · "
     "Petru & Teodor", size=13, color=MUTED)

# ---------------------------------------------------------------- 2. problem
s = new_slide("A loan or KYB file is weeks of waiting and re-checking, not one task. Raiffeisen itself described an earlier "
              "remediation as processed manually. Sources: research/07 (vendor survey), research/01.")
title(s, "Back-office work is an open loop", "Documents arrive over days; every one must agree with the others and with the policy")
stat(s, 0.6, 2.2, "95 days", "and about $2,598 per corporate KYC\nreview (2023 vendor survey)", color=REVIEW)
stat(s, 4.75, 2.2, "30 days", "a trade-register certificate stays valid,\nthen the file must ask for a new one", color=EXPIRED)
stat(s, 8.9, 2.2, "14 days", "to report a beneficial-owner mismatch\n(AMLR art. 24, from July 2027)", color=CONFLICT)
box(s, 0.6, 4.75, 12.1, 1.65)
text(s, 0.95, 4.98, 11.4, 1.3, [
    [("Today a person chases every missing paper, ", {}), ("re-checks", {"bold": True}),
     (" every expired certificate and ", {}), ("reconciles", {"bold": True}), (" owners and addresses by hand.", {})],
    [("“Procesarea desfăşurându-se manual” ", {"italic": True, "color": INK}),
     ("— Raiffeisen Bank Romania on a 14,000-client remediation (research/01)", {"color": MUTED, "size": 14})]],
    size=18, after=8)

# ---------------------------------------------------------------- 3. the real case
s = new_slide("Both fines are public (ANSPDCP 2022 and 2024, research/01). The pattern is a sequence over time, which no "
              "single-document check sees. SR-01 and SR-02 encode exactly the corrective order.")
title(s, "The risk is in the sequence, not the document", "Two real data-protection fines against Raiffeisen Bank Romania")
for i, (year, body, col) in enumerate([
        ("2022", "Credit-bureau queries without the applicant's signed consent; a client's phone number replaced with an "
                 "employee's", REVIEW),
        ("2024", "An employee changed clients' phone and email details, then applied for credit products — all without "
                 "consent", CONFLICT)]):
    y = 2.15 + i * 1.55
    dot(s, 0.6, y + 0.12, 0.34, col)
    text(s, 1.15, y, 1.4, 0.6, year, size=26, bold=True, color=col)
    text(s, 2.55, y + 0.05, 10.1, 1.2, body, size=18)
box(s, 0.6, 5.3, 12.1, 1.35, fill=SURFACE)
text(s, 0.95, 5.5, 11.4, 1.0, [
    [("OpenLoop watches the ledger: ", {"bold": True}),
     ("SR-01 flags a staff contact change followed by a credit action within 48 h with no client confirmation. "
      "SR-02: the client confirms or disputes the change in the portal; a dispute opens a card that can put the file on hold.", {})]],
    size=17)

# ---------------------------------------------------------------- 4. architecture picture
s = new_slide("One ledger per file; state is a replay of events. Code runs the pipeline; Claude plans behind a gate; the "
              "judge settles fuzzy matches; people decide on cards; both apps read the same ledger.")
title(s, "OpenLoop in one picture", "One append-only, hash-chained ledger per file — every actor reads it and writes to it")
cols = [("Events in", "documents · emails · portal uploads · staff actions · time passing", MUTED),
        ("Code pipeline", "extraction with verbatim quotes · registry lookup · hard rules HR-01…08 · timers and expiry", INK),
        ("Claude plans", "Opus 5.5 on Amazon Bedrock decides the next step over judgment-only tools", ATTN),
        ("Gate", "only actions the rules allow can run · approve credit is always DENIED", REVIEW),
        ("People decide", "short decision cards · rulings become precedents and learned checks", OK)]
for i, (head, body, col) in enumerate(cols):
    x = 0.6 + i * 2.5
    box(s, x, 2.25, 2.15, 2.55, line=LINE)
    dot(s, x + 0.25, 2.5, 0.22, col)
    text(s, x + 0.25, 2.88, 1.75, 0.5, head, size=18, bold=True, color=col if col != MUTED else INK)
    text(s, x + 0.25, 3.42, 1.75, 1.3, body, size=13, color=MUTED, spacing=1.12)
    if i < len(cols) - 1:
        arrow(s, x + 2.17, 3.52, x + 2.48, 3.52)
box(s, 0.6, 5.2, 5.95, 1.05, line=LINE)
text(s, 0.9, 5.38, 5.4, 1.1, [[("Judge  ", {"bold": True, "color": EXPIRED}),
                               ("Featherless classifier + similarity check — auto-accept only when both agree", {})]],
     size=16)
box(s, 6.75, 5.2, 5.95, 1.05, line=LINE)
text(s, 7.05, 5.38, 5.4, 1.1, [[("Two screens, one ledger  ", {"bold": True, "color": ATTN}),
                                ("the bank's work queue and the client portal, each in its own words", {})]], size=16)

# ---------------------------------------------------------------- 5. demo storyline
s = new_slide("This is the video. Each step is one click on the demo toolbar or one analyst decision.")
title(s, "One file, start to decision", "The bakery's SME loan and KYB file — the demo, in six moments")
steps = [("Day 0", "4 documents arrive; DSCR 1.4 verified; the agent asks for what's missing", ATTN),
         ("Register", "an undeclared 25% owner starts the legal clock; the address is written two ways", CONFLICT),
         ("The gate", "a client email arrives; the agent tries to ask for a new declaration — denied until a person decides", REVIEW),
         ("Decided", "two analyst clicks; corrected documents; the clock is cancelled; file complete", OK),
         ("Expiry", "the certificate expires a week later; the engine reopens the file by itself", EXPIRED),
         ("Insider", "a staff phone change, then a credit check; the client says “No, I didn't”; file on hold", CONFLICT)]
for i, (head, body, col) in enumerate(steps):
    x = 0.6 + i * 2.06
    text(s, x, 2.15, 1.9, 0.5, f"{i + 1}", size=30, bold=True, color=col)
    dot(s, x, 2.85, 0.2, col)
    if i < len(steps) - 1:
        line = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x + 0.24), Inches(2.95), Inches(x + 2.02), Inches(2.95))
        line.line.color.rgb, line.line.width = rgb(LINE), Pt(1)
    text(s, x, 3.3, 1.85, 0.5, head, size=18, bold=True)
    text(s, x, 3.8, 1.85, 2.4, body, size=14, color=MUTED, spacing=1.15)
text(s, 0.6, 6.5, 12.1, 0.5, "The credit decision stays human: the agent has no tool that approves credit.", size=16,
     color=INK, bold=True)

# ---------------------------------------------------------------- 6. agent design
s = new_slide("We first let Claude plan everything; it skipped checks and took 45-60 s per click. The hybrid measured live: "
              "3.5-13 s per click, zero fallbacks. Every step is logged with its author.")
title(s, "Agent design: a hybrid by evidence", "Code where rules are certain, Claude where judgment is needed, people where it matters")
cards = [("Code always runs", "extraction with verified quotes, registry, rules, timers, expiry — deterministic", INK),
         ("Claude plans", "next steps, email intent, escalation wording; 20 s budget per click", ATTN),
         ("Fails safe", "errors or timeouts fall back to the scripted planner; a code safety net runs after Claude", OK),
         ("Gate on every action", "allowed by the rules or denied and logged (GATE-01/02/03)", REVIEW),
         ("Humans in the loop", "cards show only what differs; rulings become precedents", EXPIRED),
         ("Learns safely", "3 similar rulings propose a check; backtested; a different person approves", CONFLICT)]
for i, (head, body, col) in enumerate(cards):
    x, y = 0.6 + (i % 3) * 4.07, 2.15 + (i // 3) * 2.15
    box(s, x, y, 3.85, 1.95, line=LINE)
    dot(s, x + 0.3, y + 0.3, 0.22, col)
    text(s, x + 0.7, y + 0.22, 3.0, 0.5, head, size=18, bold=True)
    text(s, x + 0.3, y + 0.8, 3.3, 1.1, body, size=14, color=MUTED, spacing=1.15)
text(s, 0.6, 6.6, 12.1, 0.5, "Measured live: 3.5–13 s per click, 0 fallbacks · every step logged as pipeline / claude / "
     "planner / safety-net", size=15, color=MUTED)

# ---------------------------------------------------------------- 7. proof
s = new_slide("Calibration on 150 labelled pairs: the classifier alone auto-approved 7 wrong matches (e.g. Str. Gării vs "
              "Calea Gării at p = 1.0). Requiring the similarity check too gave 0. backend/eval/calibration.json.")
title(s, "Tested, calibrated, measured", "Evidence in the repo, reproducible offline")
stat(s, 0.6, 2.1, "132", "automated tests: the full demo, every rule\npath, role leaks, tamper detection", color=INK, w=3.4)
stat(s, 0.6, 4.25, "8 / 8", "synthetic companies end on their intended\nproblem (one rule path each)", color=OK, w=3.4)
cd = CategoryChartData()
cd.categories = ["Classifier alone", "Classifier + similarity"]
cd.add_series("False approvals", (7, 0))
gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(4.4), Inches(2.0), Inches(4.4), Inches(4.3), cd)
ch = gf.chart
ch.has_legend = False
ch.has_title = True
ch.chart_title.text_frame.text = "Judge: false approvals on 150 pairs"
tp = ch.chart_title.text_frame.paragraphs[0].runs[0].font
tp.size, tp.color.rgb, tp.name, tp.bold = Pt(14), rgb(INK), FONT, False
plot = ch.plots[0]
plot.gap_width = 80
plot.has_data_labels = True
dl = plot.data_labels
dl.font.size, dl.font.color.rgb, dl.font.name, dl.font.bold = Pt(16), rgb(INK), FONT, True
dl.position = XL_LABEL_POSITION.OUTSIDE_END
for i, col in enumerate((CONFLICT, OK)):
    pt = plot.series[0].points[i]
    pt.format.fill.solid()
    pt.format.fill.fore_color.rgb = rgb(col)
for ax in (ch.category_axis, ch.value_axis):
    ax.tick_labels.font.size, ax.tick_labels.font.color.rgb, ax.tick_labels.font.name = Pt(12), rgb(MUTED), FONT
    ax.format.line.color.rgb = rgb(LINE)
ch.value_axis.has_major_gridlines = False
ch.value_axis.maximum_scale = 8
stat(s, 9.2, 2.1, "8 → 0", "false approvals of the learned address\ncheck after small-town counterexamples", color=EXPIRED, w=3.5)
stat(s, 9.2, 4.25, "390–2560", "pixel widths browser-tested, both apps,\nno console errors", color=ATTN, w=3.5)

# ---------------------------------------------------------------- 8. architecture & reproducibility
s = new_slide("SQLite for the demo; in production the same append-only table runs in PostgreSQL on the bank's own servers. "
              "Offline mode runs the whole demo with no keys and no cost.")
title(s, "Reproducible in one command", "FastAPI · SQLite event store · React bank app and client portal · Claude on Bedrock · Featherless")
box(s, 0.6, 2.15, 6.0, 1.75, fill="0B0D10", line=LINE)
text(s, 0.9, 2.4, 5.5, 1.6, ["pip install -r requirements.txt",
                             "cd web && npm ci && npm run build && cd ..",
                             "python -m uvicorn app.main:app --app-dir backend",
                             "# bank /   ·   client portal /client/"], size=14, color=OK, spacing=1.3)
rows = [("Data stays in the bank", "one append-only table; PostgreSQL on-prem in production", OK),
        ("Bring your own document AI", "extractors are adapters; quotes are verified by code", ATTN),
        ("Privacy by construction", "client views are allow-lists; staff identities never reach clients", REVIEW),
        ("All synthetic data", "documents follow real ONRC / ANAF / F20 forms; tax IDs in an unassigned range", EXPIRED)]
for i, (head, body, col) in enumerate(rows):
    y = 2.15 + i * 1.12
    dot(s, 7.0, y + 0.08, 0.2, col)
    text(s, 7.4, y, 5.3, 0.45, head, size=17, bold=True)
    text(s, 7.4, y + 0.42, 5.3, 0.6, body, size=14, color=MUTED)
text(s, 0.6, 4.25, 6.0, 1.8, [[("Offline mode ", {"bold": True}), ("(scripted planner, heuristic judge) runs the whole "
                                                                     "demo with no API keys. ", {}),
                               ("Live mode ", {"bold": True}), ("swaps in Claude and the Featherless judge with the same "
                                                              "UI and the same ledger.", {})]], size=16, spacing=1.2)

# ---------------------------------------------------------------- 9. domain
s = new_slide("All rules are code with rule IDs; the judge only settles fuzzy matches after the hard facts match.")
title(s, "Romanian rules, in code", "Two templates on one engine: SME loan and KYB check · Annual business review")
rules_ = [("HR-01", "CUI checksum"), ("HR-02", "Company active in the trade register"),
          ("HR-03", "Certificate under 30 days old"), ("HR-04", "Owners ≥25% match the register (Law 129/2019)"),
          ("HR-05", "An ID for every owner"), ("HR-06", "Same registered office everywhere"),
          ("HR-07", "Debt service coverage ≥ 1.2"), ("HR-08", "No tax debts on the ANAF certificate"),
          ("SR-01", "Staff contact change → credit action in 48 h"), ("SR-02", "Client disputes a staff change")]
for i, (rid, body) in enumerate(rules_):
    x, y = 0.6 + (i % 2) * 6.15, 2.1 + (i // 2) * 0.86
    box(s, x, y, 5.95, 0.72, line=LINE)
    text(s, x + 0.25, y + 0.18, 1.1, 0.4, rid, size=15, bold=True, color=CONFLICT if rid.startswith("SR") else ATTN)
    text(s, x + 1.4, y + 0.18, 4.4, 0.4, body, size=15)
text(s, 0.6, 6.5, 12.1, 0.5, "Judge questions JQ-01 (same address?) and JQ-02 (signatory = administrator?) decide only "
     "after the hard facts match.", size=14, color=MUTED)

# ---------------------------------------------------------------- 10. business
s = new_slide("Volume: Smart Business >100,000 active SME clients (research/01). AMLR periodic refresh from July 2027. "
              "Savings stay estimates until a shadow-mode pilot measures them.")
title(s, "Where it pays off", "The agent does the work between human decisions")
stat(s, 0.6, 2.1, "100,000+", "active Smart Business SME clients at\nRaiffeisen Romania (H1 2026)", color=ATTN, w=3.9)
stat(s, 0.6, 4.25, "1 or 5 yrs", "AMLR periodic refresh by risk, from\nJuly 2027: every SME, again and again",
     color=REVIEW, w=3.9)
box(s, 4.9, 2.1, 7.8, 4.3, line=LINE)
text(s, 5.25, 2.35, 7.2, 0.5, "Path to production", size=20, bold=True)
for i, (head, body) in enumerate([("Shadow mode", "annual business reviews with one back-office team"),
                                  ("SME onboarding", "the full loan and KYB template"),
                                  ("Same engine", "garnishments (2.25M ANAF acts in H1 2025), DORA incident clocks, "
                                                  "complaint deadlines")]):
    y = 3.05 + i * 1.05
    text(s, 5.25, y, 0.5, 0.5, f"{i + 1}", size=24, bold=True, color=OK)
    text(s, 5.8, y + 0.02, 6.6, 0.45, head, size=17, bold=True)
    text(s, 5.8, y + 0.45, 6.6, 0.5, body, size=14, color=MUTED)
text(s, 0.6, 6.6, 12.1, 0.4, "Analysts see only files that need a decision · clean files finish on their own · every "
     "action is auditable", size=15, color=MUTED)

# ---------------------------------------------------------------- 11. originality
s = new_slide("What is new compared with document-reading AI and chat assistants.")
title(s, "Why it's new")
items = [("Time is first-class", "the file is a ledger; rules run over sequences, not single documents", ATTN),
         ("An agent that cannot overstep", "a gate between plans and actions; credit decisions stay human", REVIEW),
         ("It learns from people, safely", "rulings → precedents → backtested checks with maker-checker approval", OK),
         ("One ledger, two audiences", "the bank and the client see the same file, each in their own words", EXPIRED)]
for i, (head, body, col) in enumerate(items):
    x, y = 0.6 + (i % 2) * 6.15, 1.75 + (i // 2) * 2.4
    box(s, x, y, 5.95, 2.1, line=LINE)
    dot(s, x + 0.35, y + 0.4, 0.3, col)
    text(s, x + 0.9, y + 0.33, 4.8, 0.55, head, size=22, bold=True)
    text(s, x + 0.35, y + 1.0, 5.3, 0.9, body, size=16, color=MUTED, spacing=1.15)

# ---------------------------------------------------------------- 12. ask
s = new_slide("Close: the ask is a measured pilot, not a promise.")
ledger_motif(s, 1.2, colors={2: ATTN, 5: OK, 9: OK, 13: OK})
text(s, 0.6, 2.0, 12, 1.0, "A 4-week shadow pilot", size=48, bold=True)
text(s, 0.6, 3.1, 11.5, 0.6, "on annual business reviews with one back-office team", size=24, color=INK)
for i, (head, body) in enumerate([("Plug in", "the bank's document AI and the ONRC / ANAF data feeds, on-prem"),
                                  ("Measure", "hours per file, days to complete, re-requests and audit findings"),
                                  ("Decide", "with numbers from the bank's own files")]):
    y = 4.2 + i * 0.75
    dot(s, 0.6, y + 0.1, 0.22, OK)
    text(s, 1.0, y, 11.5, 0.5, [[(head + "  ", {"bold": True}), (body, {"color": MUTED})]], size=20)
text(s, 0.6, 6.65, 12, 0.4, "OpenLoop · github.com/Petsteb/openloop-case-engine (private) · synthetic data only",
     size=13, color=MUTED)

prs.save(OUT)
print(OUT)
